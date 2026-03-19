import argparse
import csv
import html
import json
import logging
import os
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import requests
from dotenv import load_dotenv
from difflib import SequenceMatcher

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")

OPEN_QUESTION_TEXT = "Er eitthvað sem þú vilt bæta við eða segja?"
AI_FEEDBACK_SNIPPET = "AI-endurgjöfin"

load_dotenv()

INSTITUTION_URL = os.getenv("INSTITUTION_URL")
API_VERSION = os.getenv("API_VERSION", "v1")
API_TOKEN = os.getenv("API_TOKEN")
COURSE_ID = os.getenv("COURSE_ID")

BASE_URL = f"{INSTITUTION_URL}/api/{API_VERSION}" if INSTITUTION_URL else None


def validate_env_vars() -> None:
    if not API_TOKEN:
        raise ValueError("API_TOKEN is missing. Please set it in the .env file.")
    if not COURSE_ID:
        raise ValueError("COURSE_ID is missing. Please set it in the .env file.")
    if not INSTITUTION_URL:
        raise ValueError("INSTITUTION_URL is missing. Please set it in the .env file.")


def get_headers() -> dict:
    return {"Authorization": f"Bearer {API_TOKEN}"}


def _normalized_header(value: str) -> str:
    text = (value or "").strip().lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def _is_empty_value(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    return False


def _clean_tabular_value(value) -> str:
    text = str(value or "")
    text = re.sub(r"<[^>]+>", " ", text)
    text = html.unescape(text)
    return " ".join(text.split()).strip()


def _find_tabular_rows(payload):
    if isinstance(payload, list) and payload and all(isinstance(item, dict) for item in payload):
        return payload

    if isinstance(payload, dict):
        preferred_keys = [
            "rows",
            "data",
            "results",
            "items",
            "submissions",
            "quiz_submissions",
            "student_analysis",
        ]
        for key in preferred_keys:
            rows = payload.get(key)
            if isinstance(rows, list) and rows and all(isinstance(item, dict) for item in rows):
                return rows

        for value in payload.values():
            rows = _find_tabular_rows(value)
            if rows:
                return rows

    return None


def _looks_like_legacy_payload(data: dict) -> bool:
    return isinstance(data, dict) and "responses" in data and isinstance(data.get("responses"), dict)


def _extract_groups_from_row(row: dict, header_map: dict[str, str]) -> list[str]:
    group_values = []
    for key, original_key in header_map.items():
        if "group" not in key:
            continue
        value = row.get(original_key)
        if _is_empty_value(value):
            continue
        if isinstance(value, str):
            cleaned = _clean_tabular_value(value)
            parts = [part.strip() for part in re.split(r"[;,|]", cleaned) if part.strip()]
            group_values.extend(parts or [cleaned])
        else:
            group_values.append(_clean_tabular_value(value))
    return group_values


def _coerce_submitted(row: dict, header_map: dict[str, str]) -> bool:
    for key, original_key in header_map.items():
        if key in {"submitted", "submission_submitted"}:
            value = row.get(original_key)
            text = str(value).strip().lower()
            if not text:
                return False
            if text in {"false", "0", "no", "not submitted", "unsubmitted"}:
                return False
            return True
        if key in {"workflow_state", "status", "submission_status"}:
            value = row.get(original_key)
            text = str(value).strip().lower()
            return text in {"submitted", "complete", "completed", "graded"}
    return True


def _extract_name(row: dict, header_map: dict[str, str], fallback_id: str) -> str:
    preferred = [
        "student_name",
        "name",
        "student",
        "user_name",
        "full_name",
        "display_name",
        "sortable_name",
    ]
    for key in preferred:
        original_key = header_map.get(key)
        if not original_key:
            continue
        value = row.get(original_key)
        if not _is_empty_value(value):
            return _clean_tabular_value(value)
    return fallback_id


def _extract_login_id(row: dict, header_map: dict[str, str]) -> str | None:
    preferred = ["login_id", "sis_user_id", "email", "login"]
    for key in preferred:
        original_key = header_map.get(key)
        if not original_key:
            continue
        value = row.get(original_key)
        if not _is_empty_value(value):
            return _clean_tabular_value(value)
    return None


def _extract_user_id(row: dict, header_map: dict[str, str], index: int) -> str:
    preferred = [
        "user_id",
        "canvas_user_id",
        "student_id",
        "id",
        "student_canvas_id",
    ]
    for key in preferred:
        original_key = header_map.get(key)
        if not original_key:
            continue
        value = row.get(original_key)
        if not _is_empty_value(value):
            return _clean_tabular_value(value)
    return f"row_{index}"


def _question_answer_pairs_from_row(row: dict, header_map: dict[str, str]) -> list[dict]:
    metadata_prefixes = (
        "group",
        "user_id",
        "canvas_user_id",
        "student_id",
        "student_canvas_id",
        "id",
        "student_name",
        "name",
        "student",
        "user_name",
        "full_name",
        "display_name",
        "sortable_name",
        "login_id",
        "sis_user_id",
        "email",
        "login",
        "workflow_state",
        "status",
        "submission_status",
        "submitted",
        "submission_submitted",
    )

    answers = []
    for normalized_key, original_key in header_map.items():
        if normalized_key.startswith(metadata_prefixes):
            continue
        value = row.get(original_key)
        if _is_empty_value(value):
            continue
        answers.append(
            {
                "question_id": original_key,
                "question_text": original_key.strip(),
                "question_type": "tabular_import",
                "answer": _clean_tabular_value(value),
            }
        )
    return answers


def load_tabular_peer_eval_data(rows: list[dict], source_name: str) -> dict:
    responses = {}
    for idx, row in enumerate(rows, start=1):
        header_map = {_normalized_header(key): key for key in row.keys()}
        user_id = _extract_user_id(row, header_map, idx)
        name = _extract_name(row, header_map, user_id)
        answers = _question_answer_pairs_from_row(row, header_map)
        if not answers:
            continue

        response_key = str(user_id) if str(user_id) not in responses else f"{user_id}_{idx}"
        responses[response_key] = {
            "user_id": user_id,
            "name": name,
            "login_id": _extract_login_id(row, header_map),
            "groups": _extract_groups_from_row(row, header_map),
            "submitted": _coerce_submitted(row, header_map),
            "submission": {"source": source_name, "row_number": idx},
            "answers": answers,
        }

    return {
        "generated_at": datetime.now().isoformat(),
        "groups": {},
        "responses": responses,
    }


def load_peer_eval_input(input_path: Path) -> tuple[dict, str]:
    if input_path.suffix.lower() == ".csv":
        with input_path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        return load_tabular_peer_eval_data(rows, source_name="csv"), "tabular_csv"

    with input_path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)

    if _looks_like_legacy_payload(data):
        return data, "legacy_json"

    rows = _find_tabular_rows(data)
    if rows:
        return load_tabular_peer_eval_data(rows, source_name="json"), "tabular_json"

    raise ValueError(
        f"Unsupported input format in {input_path}. Expected legacy peer-eval JSON or a tabular CSV/JSON export."
    )


def fetch_group_sets() -> list[dict]:
    validate_env_vars()
    url = f"{BASE_URL}/courses/{COURSE_ID}/group_categories"
    headers = get_headers()
    group_sets = []

    while url:
        response = requests.get(url, headers=headers)
        response.raise_for_status()
        group_sets.extend(response.json())
        url = response.links.get("next", {}).get("url")

    return group_sets


def resolve_group_set_id(group_set: str, group_sets: list[dict]) -> int:
    if group_set.isdigit():
        return int(group_set)

    lowered = group_set.strip().lower()
    exact = [gs for gs in group_sets if (gs.get("name") or "").strip().lower() == lowered]
    if len(exact) == 1:
        return int(exact[0]["id"])
    if len(exact) > 1:
        raise ValueError(f"Multiple group sets match '{group_set}'. Use the numeric id.")

    partial = [gs for gs in group_sets if lowered in (gs.get("name") or "").lower()]
    if len(partial) == 1:
        return int(partial[0]["id"])
    if len(partial) > 1:
        raise ValueError(f"Multiple group sets match '{group_set}'. Use the numeric id.")

    raise ValueError(f"No group set found matching '{group_set}'.")


def fetch_group_members(group_id: int) -> list[dict]:
    url = f"{BASE_URL}/groups/{group_id}/users"
    headers = get_headers()
    members = []

    while url:
        response = requests.get(url, headers=headers)
        response.raise_for_status()
        members.extend(response.json())
        url = response.links.get("next", {}).get("url")

    cleaned = []
    for member in members:
        cleaned.append({"id": member.get("id"), "name": member.get("name")})
    return cleaned


def fetch_groups_for_set(group_set: str) -> dict:
    validate_env_vars()
    group_sets = fetch_group_sets()
    group_set_id = resolve_group_set_id(group_set, group_sets)

    url = f"{BASE_URL}/group_categories/{group_set_id}/groups"
    headers = get_headers()
    groups = {}

    while url:
        response = requests.get(url, headers=headers)
        response.raise_for_status()
        for group in response.json():
            group_id = group.get("id")
            group_name = group.get("name")
            groups[group_name] = fetch_group_members(group_id)
        url = response.links.get("next", {}).get("url")

    return groups


def get_quiz_assignment_id(quiz_id: int) -> int | None:
    validate_env_vars()
    url = f"{BASE_URL}/courses/{COURSE_ID}/quizzes/{quiz_id}"
    headers = get_headers()
    response = requests.get(url, headers=headers)
    response.raise_for_status()
    quiz_data = response.json()
    return quiz_data.get("assignment_id")


def post_grades(assignment_id: int, grade_data: dict) -> dict:
    validate_env_vars()
    url = f"{BASE_URL}/courses/{COURSE_ID}/assignments/{assignment_id}/submissions/update_grades"
    headers = get_headers()
    response = requests.post(url, headers=headers, json={"grade_data": grade_data})
    response.raise_for_status()
    return response.json()


def normalize_name(name: str) -> str:
    if not name:
        return ""
    text = unicodedata.normalize("NFKD", name)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"\s+", " ", text).strip().lower()
    return text


def override_key(group_name: str, raw_name: str) -> str:
    return f"{group_name}::{normalize_name(raw_name)}"


def load_overrides(path: Path) -> dict[str, str]:
    if not path.exists():
        return {}
    with path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)
    if isinstance(data, dict):
        return data
    return {}


def save_overrides(path: Path, overrides: dict[str, str]) -> None:
    if not overrides:
        return
    with path.open("w", encoding="utf-8") as handle:
        json.dump(overrides, handle, ensure_ascii=False, indent=2)


def prompt_for_match(
    evaluator_name: str,
    group_name: str,
    raw_name: str,
    peer_points: str | None,
    peer_reason: str | None,
    members: list[str],
) -> str | None:
    print("\nUnmatched peer name")
    print(f"Evaluator: {evaluator_name}")
    print(f"Group: {group_name}")
    print(f"Raw name: {raw_name}")
    if peer_points:
        print(f"Points: {peer_points}")
    if peer_reason:
        print(f"Reason: {peer_reason}")
    print("\nSelect the intended student:")
    for idx, member in enumerate(members, 1):
        print(f"{idx}. {member}")
    print("0. Skip / leave unmatched")

    while True:
        choice = input("Enter number: ").strip()
        if choice == "0" or choice == "":
            return None
        if choice.isdigit():
            index = int(choice)
            if 1 <= index <= len(members):
                return members[index - 1]
        print("Invalid selection. Try again.")


def prompt_for_normalize(
    evaluator_name: str,
    group_name: str,
    total: float,
    values: dict[str, float],
) -> bool:
    print("\nRow does not sum to 100.")
    print(f"Evaluator: {evaluator_name}")
    print(f"Group: {group_name}")
    print(f"Current total: {total:.2f}")
    for member, value in values.items():
        print(f"- {member}: {value}")

    while True:
        choice = input("Normalize to 100? [y/n]: ").strip().lower()
        if choice in {"y", "yes"}:
            return True
        if choice in {"n", "no", ""}:
            return False
        print("Invalid selection. Try again.")


def classify_ai_feedback(text: str) -> str:
    if not text:
        return "other"
    lowered = text.lower()
    if "alls ekki" in lowered or "lítið" in lowered:
        return "low"
    if "nokkuð" in lowered or "nokkuð" in lowered:
        return "mid"
    if "mjög" in lowered or "einstaklega" in lowered:
        return "high"
    return "other"


def _similarity_score(left: str, right: str) -> float:
    if not left or not right:
        return 0.0
    seq = SequenceMatcher(None, left, right).ratio()
    left_tokens = " ".join(sorted(left.split()))
    right_tokens = " ".join(sorted(right.split()))
    token = SequenceMatcher(None, left_tokens, right_tokens).ratio()
    return max(seq, token)


def match_name(raw_name: str, norm_map: dict[str, str], threshold: float) -> tuple[str | None, dict]:
    norm = normalize_name(raw_name)
    if not norm:
        return None, {"reason": "empty"}
    if norm in norm_map:
        return norm_map[norm], {"score": 1.0, "method": "exact"}

    # Best-effort substring match
    candidates = [(key, orig) for key, orig in norm_map.items() if norm and (norm in key or key in norm)]
    if len(candidates) == 1:
        return candidates[0][1], {"score": 0.99, "method": "substring"}

    best = None
    second = None
    for key, orig in norm_map.items():
        score = _similarity_score(norm, key)
        if best is None or score > best[0]:
            second = best
            best = (score, key, orig)
        elif second is None or score > second[0]:
            second = (score, key, orig)

    if not best:
        return None, {"reason": "no_candidates"}

    best_score, _best_key, best_orig = best
    second_score = second[0] if second else 0.0

    if best_score >= threshold and (best_score - second_score) >= 0.03:
        return best_orig, {"score": best_score, "method": "fuzzy"}

    return None, {
        "score": best_score,
        "candidate": best_orig,
        "second_score": second_score,
        "method": "ambiguous" if best_score >= threshold else "below_threshold",
    }


def parse_peer_answers(answers: list[dict]) -> tuple[list[dict], dict]:
    slots: dict[str, dict] = defaultdict(dict)
    other = {}

    for answer in answers or []:
        qtext = (answer.get("question_text") or "").strip()
        qanswer = (answer.get("answer") or "").strip()
        qtext_lower = qtext.lower()

        if "teymismeðlimur" in qtext_lower or "peer " in qtext_lower:
            match = re.search(r"(?:Teymismeðlimur|Peer)\s*(\d+)", qtext, flags=re.IGNORECASE)
            slot = match.group(1) if match else "?"
            if "nafn" in qtext_lower or "name" in qtext_lower:
                slots[slot]["name"] = qanswer
            elif "veitt stig" in qtext_lower or "points" in qtext_lower:
                slots[slot]["points"] = qanswer
            elif "athugasemdir" in qtext_lower or "justification" in qtext_lower or "comment" in qtext_lower:
                slots[slot]["reason"] = qanswer
            continue

        if qtext == OPEN_QUESTION_TEXT or "Er eitthvað sem þú vilt bæta við" in qtext:
            other["open"] = qanswer
        elif AI_FEEDBACK_SNIPPET in qtext:
            other["ai_feedback"] = qanswer

    peers = []
    for slot, entry in sorted(slots.items(), key=lambda item: int(item[0]) if item[0].isdigit() else item[0]):
        name = entry.get("name", "").strip()
        points = entry.get("points", "").strip()
        reason = entry.get("reason", "").strip()
        if name or points or reason:
            peers.append(
                {
                    "slot": slot,
                    "name": name,
                    "points": points,
                    "reason": reason,
                }
            )

    return peers, other


def build_group_maps(groups: dict) -> dict[str, dict]:
    group_maps = {}
    for group_name, members in groups.items():
        norm_map = {}
        member_ids_by_name = {}
        member_names = []
        for member in members:
            name = member.get("name", "")
            member_names.append(name)
            norm_map[normalize_name(name)] = name
            member_ids_by_name[name] = member.get("id")
        group_maps[group_name] = {
            "members": members,
            "member_names": member_names,
            "norm_map": norm_map,
            "member_ids_by_name": member_ids_by_name,
        }
    return group_maps


def count_matches(peer_names: list[str], group_map: dict, threshold: float) -> int:
    matches = 0
    for name in peer_names:
        matched, _info = match_name(name, group_map["norm_map"], threshold)
        if matched:
            matches += 1
    return matches


def choose_group(
    peer_names: list[str],
    response_groups: list[str],
    group_maps: dict[str, dict],
    threshold: float,
) -> str | None:
    if response_groups:
        candidates = [g for g in response_groups if g in group_maps]
        if not candidates:
            candidates = list(group_maps.keys())
    else:
        candidates = list(group_maps.keys())
    if not candidates:
        return None

    scored = []
    for group_name in candidates:
        group_map = group_maps[group_name]
        score = count_matches(peer_names, group_map, threshold)
        scored.append((score, len(group_map["member_names"]), group_name))

    scored.sort(key=lambda item: (-item[0], item[1], item[2]))
    return scored[0][2] if scored else None


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def write_team_csvs(team_data: dict, out_dir: Path) -> None:
    ensure_dir(out_dir)
    for team_name, data in team_data.items():
        members = data["members"]
        evaluator_map = data["evaluations"]

        safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", team_name).strip("_") or "team"
        out_path = out_dir / f"{safe_name}.csv"

        header = ["row_type", "evaluator"] + members
        with out_path.open("w", newline="", encoding="utf-8") as csv_file:
            writer = csv.writer(csv_file)
            writer.writerow(header)
            for evaluator in sorted(evaluator_map.keys()):
                points_row = ["points", evaluator]
                reason_row = ["reason", evaluator]
                eval_points = evaluator_map[evaluator]["points"]
                eval_reasons = evaluator_map[evaluator]["reasons"]
                for member in members:
                    points_row.append(eval_points.get(member, ""))
                    reason_row.append(eval_reasons.get(member, ""))
                writer.writerow(points_row)
                writer.writerow(reason_row)

        logging.info("Wrote team CSV: %s", out_path)


def parse_points_value(value: str):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        if "," in text and "." not in text:
            text = text.replace(",", ".")
        return float(text) if "." in text else int(text)
    except ValueError:
        return text


def parse_excel_points(excel_path: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError(
            "openpyxl is required for Excel input. Install with: pip install openpyxl"
        ) from exc

    workbook = load_workbook(excel_path, data_only=True)
    sheet = workbook.active

    team_data = {}
    current_team = None
    member_columns = []

    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        first = row[0]
        if isinstance(first, str) and first.strip().lower().startswith("team "):
            current_team = first.strip()[5:].strip()
            team_data[current_team] = {"members": [], "totals": defaultdict(float)}
            member_columns = []
            continue

        if not current_team:
            continue

        if isinstance(first, str) and first.strip().lower() == "row_type":
            total_idx = None
            for idx, value in enumerate(row):
                if isinstance(value, str) and value.strip().lower() == "total":
                    total_idx = idx
                    break
            if total_idx is None:
                continue
            member_columns = []
            for idx in range(2, total_idx):
                name = row[idx]
                if name:
                    member_columns.append((idx, str(name).strip()))
            team_data[current_team]["members"] = [name for _idx, name in member_columns]
            continue

        if isinstance(first, str) and first.strip().lower() == "points":
            if not member_columns:
                continue
            for idx, member_name in member_columns:
                if idx >= len(row):
                    continue
                value = parse_points_value(row[idx])
                if isinstance(value, (int, float)):
                    team_data[current_team]["totals"][member_name] += float(value)
            continue

    return team_data


def build_grade_payload_from_excel(
    excel_path: Path,
    group_maps: dict[str, dict],
    quiz_id: int,
    threshold: float,
) -> tuple[dict, list]:
    team_points = parse_excel_points(excel_path)
    grades = []
    unmatched = []
    seen_users = set()

    for team_name, data in team_points.items():
        if team_name not in group_maps:
            unmatched.append({"team": team_name, "issue": "group_not_found"})
            continue
        group_map = group_maps[team_name]
        for member_name, score in data["totals"].items():
            matched, match_info = match_name(member_name, group_map["norm_map"], threshold)
            if not matched:
                unmatched.append(
                    {
                        "team": team_name,
                        "member_name": member_name,
                        "issue": "member_not_matched",
                        "match_info": match_info,
                    }
                )
                continue

            user_id = group_map["member_ids_by_name"].get(matched)
            if not user_id:
                unmatched.append(
                    {
                        "team": team_name,
                        "member_name": member_name,
                        "issue": "missing_user_id",
                        "matched_name": matched,
                    }
                )
                continue

            if user_id in seen_users:
                unmatched.append(
                    {
                        "team": team_name,
                        "member_name": member_name,
                        "issue": "duplicate_user",
                        "user_id": user_id,
                        "matched_name": matched,
                    }
                )
                continue

            grades.append(
                {
                    "user_id": user_id,
                    "student_name": matched,
                    "team": team_name,
                    "score": int(round(score)),
                }
            )
            seen_users.add(user_id)

    grade_data = {
        str(entry["user_id"]): {"posted_grade": entry["score"]} for entry in grades
    }

    payload = {
        "generated_at": datetime.now().isoformat(),
        "quiz_id": quiz_id,
        "course_id": COURSE_ID,
        "grades": grades,
        "grade_data": grade_data,
    }

    return payload, unmatched


def write_excel_summary(team_data: dict, out_path: Path, interactive: bool = False) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
        from openpyxl.styles import PatternFill
    except ImportError as exc:
        raise ImportError(
            "openpyxl is required for Excel output. Install with: pip install openpyxl"
        ) from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Peer Evaluation"

    highlight_fill = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")

    row_cursor = 1

    for team_name in sorted(team_data.keys()):
        data = team_data[team_name]
        members = data["members"]
        evaluations = data["evaluations"]

        sheet.cell(row=row_cursor, column=1, value=f"Team {team_name}")
        row_cursor += 1

        header = ["row_type", "evaluator"] + members + ["Total"]
        for col_idx, value in enumerate(header, 1):
            sheet.cell(row=row_cursor, column=col_idx, value=value)
        header_row = row_cursor
        row_cursor += 1

        points_rows = []
        for evaluator in members:
            eval_entry = evaluations.get(
                evaluator,
                {"points": {}, "reasons": {}, "submitted": False, "auto": True},
            )
            eval_points = eval_entry.get("points", {})
            eval_reasons = eval_entry.get("reasons", {})
            submitted = bool(eval_entry.get("submitted"))
            auto = not submitted

            points_row = row_cursor
            points_rows.append(points_row)

            sheet.cell(row=points_row, column=1, value="points")
            sheet.cell(row=points_row, column=2, value=evaluator)

            row_sum = 0.0
            row_has_numeric = False
            needs_attention = False

            if auto:
                for idx, member in enumerate(members):
                    col_idx = 3 + idx
                    cell = sheet.cell(row=points_row, column=col_idx)
                    if member == evaluator:
                        cell.value = ""
                        continue
                    if len(members) > 1:
                        equal_share = 100 / (len(members) - 1)
                    else:
                        equal_share = 0
                    cell.value = equal_share
                    cell.fill = highlight_fill
                    row_sum += float(equal_share)
                    row_has_numeric = True
                needs_attention = True
            else:
                values_by_member: dict[str, float | str | None] = {}
                numeric_values: dict[str, float] = {}
                for member in members:
                    if member == evaluator:
                        continue
                    raw_value = eval_points.get(member, "")
                    parsed_value = parse_points_value(raw_value)
                    values_by_member[member] = parsed_value
                    if isinstance(parsed_value, (int, float)):
                        numeric_values[member] = float(parsed_value)
                    else:
                        needs_attention = True

                row_has_numeric = bool(numeric_values)
                row_sum = sum(numeric_values.values()) if numeric_values else 0.0

                normalized = False
                if (
                    interactive
                    and not needs_attention
                    and row_has_numeric
                    and abs(row_sum - 100) > 0.01
                ):
                    if prompt_for_normalize(evaluator, team_name, row_sum, numeric_values):
                        scale = 100 / row_sum if row_sum else 0
                        members_to_scale = [m for m in members if m != evaluator]
                        normalized_values = {}
                        running_total = 0.0
                        for member in members_to_scale[:-1]:
                            value = numeric_values.get(member, 0.0) * scale
                            value = round(value, 2)
                            normalized_values[member] = value
                            running_total += value
                        last_member = members_to_scale[-1] if members_to_scale else None
                        if last_member:
                            normalized_values[last_member] = round(100 - running_total, 2)
                        for member, value in normalized_values.items():
                            values_by_member[member] = value
                        row_sum = 100.0
                        normalized = True
                        needs_attention = True

                for idx, member in enumerate(members):
                    col_idx = 3 + idx
                    cell = sheet.cell(row=points_row, column=col_idx)
                    if member == evaluator:
                        cell.value = ""
                        continue
                    value = values_by_member.get(member, "")
                    cell.value = value
                    if normalized:
                        cell.fill = highlight_fill
                    elif value == "" or value is None or not isinstance(value, (int, float)):
                        cell.fill = highlight_fill

            total_col = 3 + len(members)
            start_col_letter = get_column_letter(3)
            end_col_letter = get_column_letter(total_col - 1)
            total_cell = sheet.cell(
                row=points_row,
                column=total_col,
                value=f"=SUM({start_col_letter}{points_row}:{end_col_letter}{points_row})",
            )
            if needs_attention or (not row_has_numeric) or (abs(row_sum - 100) > 0.01):
                total_cell.fill = highlight_fill

            reason_row = row_cursor + 1
            sheet.cell(row=reason_row, column=1, value="reason")
            sheet.cell(row=reason_row, column=2, value=evaluator)
            for idx, member in enumerate(members):
                col_idx = 3 + idx
                sheet.cell(row=reason_row, column=col_idx, value=eval_reasons.get(member, ""))

            row_cursor += 2

        total_row = row_cursor
        sheet.cell(row=total_row, column=1, value="Total")

        total_col = 3 + len(members)
        if points_rows:
            start_row = header_row + 1
            end_row = total_row - 1

            total_col_letter = get_column_letter(total_col)
            sheet.cell(
                row=total_row,
                column=2,
                value=f"=SUM({total_col_letter}{start_row}:{total_col_letter}{end_row})",
            )

            for idx, _member in enumerate(members):
                col_idx = 3 + idx
                col_letter = get_column_letter(col_idx)
                sheet.cell(
                    row=total_row,
                    column=col_idx,
                    value=f"=SUM({col_letter}{start_row}:{col_letter}{end_row})",
                )

            sheet.cell(
                row=total_row,
                column=total_col,
                value=f"=SUM({total_col_letter}{start_row}:{total_col_letter}{end_row})",
            )

        row_cursor += 2

    workbook.save(out_path)
    logging.info("Wrote Excel summary: %s", out_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Compile peer evaluation quiz data into team CSVs and summaries.")
    parser.add_argument(
        "input_json",
        nargs="?",
        default="peer_eval_20214.json",
        help="Peer evaluation input file. Supports legacy JSON plus tabular CSV/JSON exports (default: peer_eval_20214.json).",
    )
    parser.add_argument(
        "--list-group-sets",
        action="store_true",
        help="List available Canvas group sets (group categories) and exit.",
    )
    parser.add_argument(
        "--group-set",
        default=None,
        help="Group set name or id to use for team membership (overrides JSON groups).",
    )
    parser.add_argument(
        "--out-dir",
        default=None,
        help="Output directory (default: same directory as input).",
    )
    parser.add_argument(
        "--teams-dir",
        default="teams",
        help="Subdirectory for team CSV outputs (default: teams).",
    )
    parser.add_argument(
        "--open-questions-out",
        default="open_questions.txt",
        help="Output text file for open question answers.",
    )
    parser.add_argument(
        "--excel-out",
        default="peer_eval_summary.xlsx",
        help="Output Excel workbook (default: peer_eval_summary.xlsx).",
    )
    parser.add_argument(
        "--fuzzy-threshold",
        type=float,
        default=0.84,
        help="Fuzzy match threshold for peer names (default: 0.84).",
    )
    parser.add_argument(
        "--interactive",
        action="store_true",
        help="Prompt for manual name matching when fuzzy matching fails.",
    )
    parser.add_argument(
        "--overrides",
        default=None,
        help="Path to manual overrides JSON (default: manual_overrides.json in output dir).",
    )
    parser.add_argument(
        "--grades-from-excel",
        action="store_true",
        help="Generate Canvas grade payload from corrected Excel sheet.",
    )
    parser.add_argument(
        "--quiz-id",
        type=int,
        default=None,
        help="Quiz ID for grade payload generation.",
    )
    parser.add_argument(
        "--excel",
        default=None,
        help="Excel file to read grades from (peer_eval_summary.xlsx).",
    )
    parser.add_argument(
        "--grades-out",
        default=None,
        help="Output JSON file for Canvas grade payload.",
    )
    parser.add_argument(
        "--post-grades",
        action="store_true",
        help="Post grades to Canvas after generating or loading a payload.",
    )
    parser.add_argument(
        "--payload",
        default=None,
        help="Existing payload JSON to post (used with --post-grades).",
    )

    args = parser.parse_args()

    if args.list_group_sets:
        group_sets = fetch_group_sets()
        if not group_sets:
            logging.info("No group sets found.")
        else:
            logging.info("Available group sets:")
            for gs in group_sets:
                logging.info("  %s: %s", gs.get("id"), gs.get("name"))
        return

    if args.grades_from_excel:
        if not args.quiz_id:
            raise ValueError("--quiz-id is required with --grades-from-excel.")
        if not args.group_set:
            raise ValueError("--group-set is required with --grades-from-excel.")
        if not args.excel:
            raise ValueError("--excel is required with --grades-from-excel.")

        excel_path = Path(args.excel)
        if not excel_path.exists():
            raise FileNotFoundError(f"Excel file not found: {excel_path}")

        out_dir = Path(args.out_dir) if args.out_dir else excel_path.parent
        ensure_dir(out_dir)
        grades_out = (
            Path(args.grades_out)
            if args.grades_out
            else out_dir / f"quiz_{args.quiz_id}_grades_payload.json"
        )

        groups = fetch_groups_for_set(args.group_set)
        group_maps = build_group_maps(groups)
        payload, unmatched = build_grade_payload_from_excel(
            excel_path,
            group_maps,
            args.quiz_id,
            args.fuzzy_threshold,
        )

        assignment_id = get_quiz_assignment_id(args.quiz_id)
        payload["assignment_id"] = assignment_id
        if assignment_id:
            payload["endpoint"] = (
                f"/api/v1/courses/{COURSE_ID}/assignments/{assignment_id}/submissions/update_grades"
            )

        with open(grades_out, "w", encoding="utf-8") as out_file:
            json.dump(payload, out_file, ensure_ascii=False, indent=2)
        logging.info("Wrote grade payload: %s", grades_out)

        if unmatched:
            unmatched_path = out_dir / "grades_unmatched.json"
            with unmatched_path.open("w", encoding="utf-8") as out_file:
                json.dump(unmatched, out_file, ensure_ascii=False, indent=2)
            logging.warning("Unmatched grade entries written to %s", unmatched_path)

        if args.post_grades:
            if not assignment_id:
                raise ValueError("Quiz has no assignment_id; cannot post grades.")
            result = post_grades(assignment_id, payload["grade_data"])
            logging.info("Posted grades. Canvas response keys: %s", ", ".join(result.keys()))
        return

    if args.post_grades:
        if not args.payload:
            raise ValueError("--payload is required with --post-grades when not using --grades-from-excel.")

        payload_path = Path(args.payload)
        if not payload_path.exists():
            raise FileNotFoundError(f"Payload not found: {payload_path}")

        with payload_path.open("r", encoding="utf-8") as handle:
            payload = json.load(handle)

        assignment_id = payload.get("assignment_id")
        if not assignment_id:
            if not args.quiz_id:
                raise ValueError("--quiz-id is required if payload lacks assignment_id.")
            assignment_id = get_quiz_assignment_id(args.quiz_id)
        if not assignment_id:
            raise ValueError("Quiz has no assignment_id; cannot post grades.")

        grade_data = payload.get("grade_data", {})
        if not grade_data:
            raise ValueError("Payload missing grade_data.")

        result = post_grades(int(assignment_id), grade_data)
        logging.info("Posted grades. Canvas response keys: %s", ", ".join(result.keys()))
        return

    input_path = Path(args.input_json)
    if not input_path.exists():
        raise FileNotFoundError(f"Input JSON not found: {input_path}")

    out_dir = Path(args.out_dir) if args.out_dir else input_path.parent
    ensure_dir(out_dir)
    teams_dir = out_dir / args.teams_dir
    overrides_path = Path(args.overrides) if args.overrides else out_dir / "manual_overrides.json"
    overrides = load_overrides(overrides_path) if args.interactive else {}

    data, input_format = load_peer_eval_input(input_path)
    logging.info("Loaded peer evaluation input as %s.", input_format)

    if args.group_set:
        groups = fetch_groups_for_set(args.group_set)
        logging.info("Using groups from group set: %s", args.group_set)
    else:
        groups = data.get("groups", {})
        logging.info("Using groups from input JSON.")
        if not groups:
            raise ValueError(
                "Input does not include group membership data. Use --group-set <name-or-id> with New Quiz tabular exports."
            )
    responses = data.get("responses", {})

    group_maps = build_group_maps(groups)

    team_data = {}
    unmatched = []
    open_answers = []
    ai_feedback_counts = Counter()

    for response in responses.values():
        user_id = str(response.get("user_id"))
        name = response.get("name") or user_id
        response_groups = response.get("groups") or []
        answers = response.get("answers") or []

        peers, other = parse_peer_answers(answers)
        peer_names = [peer["name"] for peer in peers if peer.get("name")]

        group_name = choose_group(peer_names, response_groups, group_maps, args.fuzzy_threshold)
        if not group_name:
            unmatched.append({"user_id": user_id, "name": name, "issue": "group_not_found"})
            continue

        if group_name not in team_data:
            team_data[group_name] = {
                "members": group_maps[group_name]["member_names"],
                "evaluations": {},
            }

        eval_entry = team_data[group_name]["evaluations"].setdefault(
            name,
            {
                "points": {},
                "reasons": {},
                "user_id": user_id,
                "submitted": bool(response.get("submitted")),
            },
        )

        for peer in peers:
            raw_name = peer.get("name", "")
            override = overrides.get(override_key(group_name, raw_name))
            if override and override in group_maps[group_name]["member_names"]:
                matched = override
                match_info = {"method": "override"}
            else:
                matched, match_info = match_name(
                    raw_name, group_maps[group_name]["norm_map"], args.fuzzy_threshold
                )

            if not matched:
                if args.interactive:
                    choice = prompt_for_match(
                        name,
                        group_name,
                        raw_name,
                        peer.get("points"),
                        peer.get("reason"),
                        group_maps[group_name]["member_names"],
                    )
                    if choice:
                        overrides[override_key(group_name, raw_name)] = choice
                        matched = choice
                        match_info = {"method": "interactive"}

                if not matched:
                    unmatched.append(
                        {
                            "user_id": user_id,
                            "name": name,
                            "issue": "unmatched_peer_name",
                            "peer_name": raw_name,
                            "group": group_name,
                            "match_info": match_info,
                            "peer_points": peer.get("points"),
                            "peer_reason": peer.get("reason"),
                            "peer_slot": peer.get("slot"),
                        }
                    )
                    continue

            eval_entry["points"][matched] = peer.get("points", "")
            eval_entry["reasons"][matched] = peer.get("reason", "")

        open_text = other.get("open")
        if open_text:
            open_answers.append(
                {
                    "group": group_name,
                    "name": name,
                    "user_id": user_id,
                    "answer": open_text,
                }
            )

        ai_feedback = other.get("ai_feedback")
        if ai_feedback:
            ai_feedback_counts[ai_feedback] += 1

    if open_answers:
        open_path = out_dir / args.open_questions_out
        with open_path.open("w", encoding="utf-8") as out_file:
            out_file.write(f"Open questions exported {datetime.now().isoformat()}\n\n")
            for entry in open_answers:
                out_file.write(f"Group: {entry['group']}\n")
                out_file.write(f"Student: {entry['name']} ({entry['user_id']})\n")
                out_file.write(entry["answer"].strip() + "\n")
                out_file.write("\n---\n\n")
        logging.info("Wrote open question responses: %s", open_path)
    else:
        logging.info("No open question responses found.")

    write_team_csvs(team_data, teams_dir)
    write_excel_summary(team_data, out_dir / args.excel_out, interactive=args.interactive)

    total_ai = sum(ai_feedback_counts.values())
    if total_ai:
        sentiment_counts = Counter()
        for answer, count in ai_feedback_counts.items():
            sentiment_counts[classify_ai_feedback(answer)] += count

        logging.info("AI feedback summary (n=%s):", total_ai)
        most_common = ai_feedback_counts.most_common(1)[0]
        logging.info(
            "  Most common: %s (%s, %.1f%%)",
            most_common[0],
            most_common[1],
            (most_common[1] / total_ai) * 100,
        )

        for label in ["high", "mid", "low", "other"]:
            count = sentiment_counts.get(label, 0)
            if count:
                logging.info(
                    "  %s: %s (%.1f%%)",
                    label,
                    count,
                    (count / total_ai) * 100,
                )

        logging.info("AI feedback counts:")
        for answer, count in ai_feedback_counts.most_common():
            pct = (count / total_ai) * 100
            logging.info("  %s: %s (%.1f%%)", answer, count, pct)
    else:
        logging.info("No AI feedback responses found.")

    if unmatched:
        unmatched_path = out_dir / "unmatched_names.json"
        with unmatched_path.open("w", encoding="utf-8") as out_file:
            json.dump(unmatched, out_file, ensure_ascii=False, indent=2)
        logging.warning("Unmatched names written to %s", unmatched_path)

    if args.interactive:
        save_overrides(overrides_path, overrides)


if __name__ == "__main__":
    main()
